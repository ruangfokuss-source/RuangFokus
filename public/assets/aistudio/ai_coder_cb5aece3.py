#!/usr/bin/env python3
"""
ai_coder.py — Asisten Rekayasa Perangkat Lunak Agentic & Analisis Kode (v2).

Instalasi:
    pip install openai anthropic google-genai python-dotenv rich python-docx

Konfigurasi (.env atau environment variables):
    PROVIDER=anthropic           # openai | anthropic | gemini
    OPENAI_API_KEY=sk-...
    ANTHROPIC_API_KEY=sk-ant-...
    GEMINI_API_KEY=...
    MODEL=                       # opsional; kosongkan untuk default
    THINK=medium                 # off | low | medium | high  (deep reasoning)
    SANDBOX=local                # local | docker  (untuk tool run_python)
    DOCKER_IMAGE=python:3.12-slim
    CONTEXT_CHARS=400000         # ambang kompaksi riwayat (karakter)
    MAX_AGENT_STEPS=15

Fitur utama:
    * Streaming nyata + agentic tool-loop di ketiga provider
    * Deep thinking (Anthropic extended thinking, OpenAI reasoning_effort, Gemini thinking_budget)
    * Multimodal: gambar & PDF (semua provider), audio & video (Gemini)
    * Tools: baca/cari kode, tulis/edit berkas (dengan diff + persetujuan + /undo),
      code interpreter (run_python, opsional Docker tanpa jaringan), run_command, create_report (md/docx)
    * Kompaksi konteks otomatis lewat ringkasan (bukan sekadar membuang pesan lama)
    * Prompt caching (Anthropic), retry exponential backoff, sesi SQLite
"""

from __future__ import annotations

import base64
import difflib
import fnmatch
import json
import mimetypes
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import time
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any, Callable, Generator

from dotenv import load_dotenv
from rich.console import Console, RenderableType
from rich.live import Live
from rich.markdown import Markdown
from rich.panel import Panel
from rich.syntax import Syntax
from rich.table import Table

load_dotenv()
console = Console()

# ---------------------------------------------------------------------------
# Konfigurasi & Konstanta
# ---------------------------------------------------------------------------

APP_DIR = Path.home() / ".ai_coder"
DB_PATH = APP_DIR / "sessions.db"

MAX_HISTORY_CHARS = int(os.getenv("CONTEXT_CHARS", "400000"))
MAX_TOOL_OUTPUT = 30_000
MAX_AGENT_STEPS = int(os.getenv("MAX_AGENT_STEPS", "15"))
MAX_OUTPUT_TOKENS = 16_000
MAX_FILE_BYTES = 2_000_000
KEEP_RECENT_MESSAGES = 5          # selalu ganjil agar pesan pertama yang dipertahankan = user

SANDBOX = os.getenv("SANDBOX", "local").lower()
DOCKER_IMAGE = os.getenv("DOCKER_IMAGE", "python:3.12-slim")

THINK_LEVELS = ("off", "low", "medium", "high")
ANTHROPIC_BUDGET = {"off": 0, "low": 4_000, "medium": 10_000, "high": 24_000}
GEMINI_BUDGET = {"off": 128, "low": 2_048, "medium": 8_192, "high": 24_576}

DEFAULT_MODELS = {
    "openai": "gpt-5",
    "anthropic": "claude-sonnet-5-5",
    "gemini": "gemini-2.5-pro",
}

PROVIDER_MEDIA = {
    "openai": {"image", "pdf"},
    "anthropic": {"image", "pdf"},
    "gemini": {"image", "pdf", "audio", "video"},
}

IGNORED_DIRS = {
    ".git", "node_modules", "__pycache__", ".venv", "venv",
    "dist", "build", ".next", ".idea", ".vscode", "target", "vendor", "ai_outputs",
}

CODE_EXTENSIONS = {
    ".py", ".js", ".ts", ".jsx", ".tsx", ".java", ".go", ".rs", ".rb",
    ".php", ".c", ".cpp", ".h", ".hpp", ".cs", ".kt", ".swift", ".sql",
    ".sh", ".html", ".css", ".scss", ".json", ".yaml", ".yml", ".toml",
    ".md", ".dockerfile", ".tf", ".lua", ".txt", ".ini", ".cfg", ".env.example",
}

SYSTEM_PROMPT = """Anda adalah Principal Software Engineer, analis keamanan kode, dan analis data tingkat lanjut.
Anda bekerja langsung di workspace proyek pengguna melalui tools:
- `list_project_files`, `search_in_files`, `read_file_lines`: inspeksi kode (baca-saja).
- `write_file`, `replace_in_file`: mengubah berkas (pengguna akan melihat diff dan menyetujui).
- `run_python`: code interpreter terisolasi untuk perhitungan, analisis data, prototipe, dan verifikasi logika.
  Simpan berkas hasil (grafik, CSV) ke folder pada variabel lingkungan AI_OUTPUT_DIR.
- `run_command`: menjalankan perintah shell (tes, linter, build); selalu butuh persetujuan pengguna.
- `create_report`: membuat laporan terstruktur (Markdown atau Word .docx).

Aturan kerja:
1. Jangan menebak isi kode atau data. Periksa dengan tools terlebih dahulu; untuk berkas besar gunakan `search_in_files`, lalu baca bertahap.
2. Semua perhitungan non-trivial (angka, statistik, keuangan, kompleksitas) harus diverifikasi dengan `run_python`, bukan dihitung "di kepala".
3. Saat menganalisis: sebutkan berkas dan nomor baris, jelaskan akar masalah (bug, race condition, celah keamanan, kompleksitas berlebih), dan berikan perbaikan yang lengkap dan siap pakai.
4. Untuk perubahan kode: lakukan perubahan minimal dengan `replace_in_file`, lalu jalankan tes/linter dengan `run_command` bila tersedia.
5. Untuk tugas kompleks: susun rencana singkat, kerjakan bertahap, verifikasi hasil, lalu laporkan.
6. Jika informasi tidak cukup atau tidak pasti, katakan secara eksplisit. Jangan mengarang fakta, kutipan, atau nomor baris.
7. Jawab tajam, teknis, dan langsung ke inti; gunakan bahasa yang sama dengan pengguna.
"""


# ---------------------------------------------------------------------------
# Utilitas
# ---------------------------------------------------------------------------

def with_retry(fn: Callable[[], Any], retries: int = 4, base_delay: float = 1.5) -> Any:
    """Jalankan pemanggilan API dengan exponential backoff untuk error transien."""
    for attempt in range(retries):
        try:
            return fn()
        except Exception as exc:
            msg = str(exc).lower()
            transient = any(
                k in msg
                for k in ("rate limit", "429", "500", "502", "503", "504", "529", "overloaded", "timeout", "connection")
            )
            if not transient or attempt == retries - 1:
                raise
            wait = base_delay * (2 ** attempt)
            console.print(f"[yellow]⏳ Limit/gangguan API. Mencoba ulang dalam {wait:.1f} dtk...[/yellow]")
            time.sleep(wait)


def truncate(text: str, limit: int = MAX_TOOL_OUTPUT) -> str:
    if len(text) <= limit:
        return text
    return text[:limit] + f"\n... [dipotong, {len(text) - limit} karakter tidak ditampilkan]"


@dataclass
class Attachment:
    path: Path
    mime: str
    kind: str  # image | pdf | audio | video

    @property
    def b64(self) -> str:
        return base64.b64encode(self.path.read_bytes()).decode()


def load_attachment(path_str: str, base: Path) -> Attachment:
    p = Path(path_str).expanduser()
    if not p.is_absolute():
        p = base / p
    p = p.resolve()
    if not p.is_file():
        raise FileNotFoundError(f"Berkas tidak ditemukan: {p}")
    mime = mimetypes.guess_type(p.name)[0] or "application/octet-stream"
    if mime.startswith("image/"):
        kind = "image"
    elif mime == "application/pdf":
        kind = "pdf"
    elif mime.startswith("audio/"):
        kind = "audio"
    elif mime.startswith("video/"):
        kind = "video"
    else:
        raise ValueError(f"Format lampiran tidak didukung ({mime}). Untuk berkas teks/kode gunakan /file.")
    return Attachment(p, mime, kind)


# ---------------------------------------------------------------------------
# Persetujuan Pengguna (untuk tool yang mengubah sistem)
# ---------------------------------------------------------------------------

class Approvals:
    """Meminta konfirmasi sebelum tool yang berdampak (tulis berkas, eksekusi kode/perintah)."""

    def __init__(self) -> None:
        self.auto = False
        self.before: Callable[[], None] = lambda: None   # dipakai untuk menjeda Live UI
        self.after: Callable[[], None] = lambda: None

    def ask(self, title: str, preview: RenderableType, force: bool = False) -> bool:
        if self.auto and not force:
            return True
        self.before()
        try:
            console.print(Panel(preview, title=title, border_style="yellow"))
            hint = "y/N" if force else "y/N/a=setujui semua perubahan berkas"
            ans = console.input(f"[yellow]Izinkan? [{hint}] > [/yellow]").strip().lower()
        finally:
            self.after()
        if ans == "a" and not force:
            self.auto = True
            return True
        return ans in ("y", "yes", "ya")


# ---------------------------------------------------------------------------
# Penyimpanan Sesi (SQLite)
# ---------------------------------------------------------------------------

class SessionStore:
    def __init__(self, db_path: Path = DB_PATH):
        db_path.parent.mkdir(parents=True, exist_ok=True)
        self.conn = sqlite3.connect(db_path)
        self.conn.execute("PRAGMA journal_mode=WAL")
        self.conn.execute(
            """
            CREATE TABLE IF NOT EXISTS messages (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                session_id TEXT NOT NULL,
                role TEXT NOT NULL,
                content TEXT NOT NULL,
                created_at REAL NOT NULL
            )
            """
        )
        self.conn.execute("CREATE INDEX IF NOT EXISTS idx_session ON messages(session_id, id)")
        self.conn.commit()

    def save_message(self, session_id: str, role: str, content: str) -> None:
        self.conn.execute(
            "INSERT INTO messages (session_id, role, content, created_at) VALUES (?, ?, ?, ?)",
            (session_id, role, content, time.time()),
        )
        self.conn.commit()

    def replace_session(self, session_id: str, history: list[dict[str, str]]) -> None:
        """Tulis ulang sesi (dipakai setelah kompaksi konteks)."""
        self.conn.execute("DELETE FROM messages WHERE session_id = ?", (session_id,))
        now = time.time()
        self.conn.executemany(
            "INSERT INTO messages (session_id, role, content, created_at) VALUES (?, ?, ?, ?)",
            [(session_id, m["role"], m["content"], now + i * 1e-3) for i, m in enumerate(history)],
        )
        self.conn.commit()

    def load_session(self, session_id: str) -> list[dict[str, str]]:
        cur = self.conn.execute(
            "SELECT role, content FROM messages WHERE session_id = ? ORDER BY id ASC", (session_id,)
        )
        return [{"role": r, "content": c} for r, c in cur.fetchall()]

    def clear_session(self, session_id: str) -> None:
        self.conn.execute("DELETE FROM messages WHERE session_id = ?", (session_id,))
        self.conn.commit()

    def list_sessions(self) -> list[tuple[str, int, str]]:
        cur = self.conn.execute(
            """
            SELECT session_id, COUNT(*), datetime(MAX(created_at), 'unixepoch', 'localtime')
            FROM messages GROUP BY session_id ORDER BY MAX(created_at) DESC
            """
        )
        return cur.fetchall()


# ---------------------------------------------------------------------------
# Workspace Tools
# ---------------------------------------------------------------------------

class WorkspaceTools:
    """Tools yang dapat dipanggil AI. Semua akses file dibatasi di dalam root workspace."""

    def __init__(self, root: Path, approvals: Approvals):
        self.root = root.resolve()
        self.approvals = approvals
        self.undo_stack: list[tuple[Path, str | None]] = []
        self._handlers: dict[str, Callable[..., str]] = {
            "list_project_files": self.list_project_files,
            "search_in_files": self.search_in_files,
            "read_file_lines": self.read_file_lines,
            "write_file": self.write_file,
            "replace_in_file": self.replace_in_file,
            "run_python": self.run_python,
            "run_command": self.run_command,
            "create_report": self.create_report,
        }

    # -- util -----------------------------------------------------------------

    def set_root(self, new_root: Path) -> None:
        resolved = new_root.expanduser().resolve()
        if not resolved.exists():
            raise FileNotFoundError(f"Path tidak ditemukan: {resolved}")
        self.root = resolved if resolved.is_dir() else resolved.parent

    def _safe_path(self, rel_path: str) -> Path:
        target = (self.root / rel_path).resolve()
        if not target.is_relative_to(self.root):
            raise PermissionError(f"Akses ditolak di luar workspace: {rel_path}")
        return target

    @staticmethod
    def _is_code(p: Path) -> bool:
        return p.suffix.lower() in CODE_EXTENSIONS or p.name in ("Dockerfile", "Makefile")

    def _iter_files(self, base: Path, glob: str = "*"):
        for dirpath, dirnames, filenames in os.walk(base):
            dirnames[:] = sorted(d for d in dirnames if d not in IGNORED_DIRS)
            for name in sorted(filenames):
                p = Path(dirpath) / name
                if glob != "*" and not fnmatch.fnmatch(name, glob):
                    continue
                if glob == "*" and not self._is_code(p):
                    continue
                yield p

    def execute(self, name: str, args: dict[str, Any]) -> str:
        handler = self._handlers.get(name)
        if handler is None:
            return f"Error: Tool '{name}' tidak dikenal."
        try:
            return truncate(str(handler(**args)))
        except TypeError as e:
            return f"Error argumen tool: {e}"
        except Exception as e:  # noqa: BLE001 — hasil error dikembalikan ke model agar bisa memperbaiki
            return f"Error: {type(e).__name__}: {e}"

    # -- tools baca-saja ------------------------------------------------------

    def list_project_files(self, subdir: str = ".", max_results: int = 300) -> str:
        base = self._safe_path(subdir)
        if not base.is_dir():
            return f"Error: Direktori '{subdir}' tidak ditemukan."
        files: list[str] = []
        for p in self._iter_files(base):
            files.append(f"{p.relative_to(self.root).as_posix()} ({p.stat().st_size / 1024:.1f} KB)")
            if len(files) >= int(max_results):
                files.append(f"... [dipotong pada {max_results} berkas]")
                break
        return "\n".join(files) or "Tidak ada berkas kode ditemukan."

    def read_file_lines(self, filepath: str, start_line: int = 1, end_line: int = 400) -> str:
        target = self._safe_path(filepath)
        if not target.is_file():
            return f"Error: Berkas '{filepath}' tidak ditemukan."
        lines = target.read_text(encoding="utf-8", errors="replace").splitlines()
        total = len(lines)
        start = max(1, int(start_line))
        end = min(total, max(start, int(end_line)))
        body = "\n".join(f"{i:5d} | {lines[i - 1]}" for i in range(start, end + 1))
        return f"Berkas: {filepath} (Baris {start}-{end} dari total {total})\n{body}"

    def search_in_files(self, pattern: str, file_pattern: str = "*", context: int = 0) -> str:
        try:
            regex = re.compile(pattern, re.IGNORECASE)
        except re.error as e:
            return f"Regex tidak valid: {e}"
        context = max(0, min(int(context), 5))
        out: list[str] = []
        hits = 0
        for p in self._iter_files(self.root, file_pattern):
            try:
                if p.stat().st_size > MAX_FILE_BYTES:
                    continue
                lines = p.read_text(encoding="utf-8", errors="ignore").splitlines()
            except OSError:
                continue
            rel = p.relative_to(self.root).as_posix()
            for idx, line in enumerate(lines, start=1):
                if not regex.search(line):
                    continue
                if context:
                    lo, hi = max(1, idx - context), min(len(lines), idx + context)
                    out.append(f"--- {rel}:{idx}")
                    out.extend(f"{n:5d}{'>' if n == idx else ' '}| {lines[n - 1][:200]}" for n in range(lo, hi + 1))
                else:
                    out.append(f"{rel}:{idx}: {line.strip()[:160]}")
                hits += 1
                if hits >= 100:
                    out.append("... [dibatasi 100 kecocokan; persempit pola pencarian]")
                    return "\n".join(out)
        return "\n".join(out) or f"Tidak ditemukan kecocokan untuk '{pattern}'."

    # -- tools yang mengubah berkas ------------------------------------------

    def _apply_write(self, target: Path, rel: str, new_content: str) -> str:
        old = target.read_text(encoding="utf-8", errors="replace") if target.exists() else None
        diff = "".join(
            difflib.unified_diff(
                (old or "").splitlines(True), new_content.splitlines(True), f"a/{rel}", f"b/{rel}"
            )
        )
        if old == new_content:
            return "Tidak ada perubahan (isi identik)."
        shown = "\n".join(diff.splitlines()[:200]) or "(berkas baru/kosong)"
        if not self.approvals.ask(f"{'Ubah' if old is not None else 'Buat'}: {rel}", Syntax(shown, "diff")):
            return "DITOLAK oleh pengguna. Jangan ulangi tanpa mengubah pendekatan; tanyakan pengguna bila perlu."
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(new_content, encoding="utf-8")
        self.undo_stack.append((target, old))
        return f"OK: {rel} ditulis ({len(new_content)} karakter)."

    def write_file(self, filepath: str, content: str) -> str:
        return self._apply_write(self._safe_path(filepath), filepath, content)

    def replace_in_file(self, filepath: str, old_str: str, new_str: str) -> str:
        target = self._safe_path(filepath)
        if not target.is_file():
            return f"Error: Berkas '{filepath}' tidak ditemukan."
        if not old_str:
            return "Error: old_str tidak boleh kosong."
        text = target.read_text(encoding="utf-8")
        n = text.count(old_str)
        if n != 1:
            return f"Error: old_str ditemukan {n} kali; harus tepat 1. Tambahkan konteks di sekitarnya."
        return self._apply_write(target, filepath, text.replace(old_str, new_str, 1))

    def undo(self) -> str:
        if not self.undo_stack:
            return "Tidak ada perubahan yang bisa di-undo."
        path, old = self.undo_stack.pop()
        if old is None:
            path.unlink(missing_ok=True)
            return f"Berkas baru dihapus: {path.relative_to(self.root).as_posix()}"
        path.write_text(old, encoding="utf-8")
        return f"Dikembalikan: {path.relative_to(self.root).as_posix()}"

    # -- code interpreter & shell --------------------------------------------

    @staticmethod
    def _clean_env() -> dict[str, str]:
        return {
            k: v for k, v in os.environ.items()
            if not any(s in k.upper() for s in ("KEY", "TOKEN", "SECRET", "PASSWORD"))
        }

    def run_python(self, code: str, timeout: int = 30) -> str:
        timeout = max(1, min(int(timeout), 180))
        use_docker = SANDBOX == "docker" and shutil.which("docker") is not None
        out_dir = self.root / "ai_outputs"

        if not use_docker:
            preview = Syntax(code, "python", line_numbers=True)
            if not self.approvals.ask("Jalankan kode Python (LOKAL, tanpa isolasi penuh)", preview, force=True):
                return "DITOLAK oleh pengguna."
        out_dir.mkdir(exist_ok=True)

        with tempfile.TemporaryDirectory() as tmp:
            (Path(tmp) / "main.py").write_text(code, encoding="utf-8")
            if use_docker:
                cmd = [
                    "docker", "run", "--rm", "--network", "none", "--memory", "1g", "--cpus", "1",
                    "-v", f"{tmp}:/work", "-v", f"{out_dir}:/out", "-e", "AI_OUTPUT_DIR=/out",
                    "-w", "/work", DOCKER_IMAGE, "python", "main.py",
                ]
                env = None
            else:
                cmd = [sys.executable, "-I", "main.py"]
                env = {**self._clean_env(), "AI_OUTPUT_DIR": str(out_dir), "PYTHONIOENCODING": "utf-8"}
            try:
                p = subprocess.run(cmd, cwd=tmp, capture_output=True, text=True, timeout=timeout, env=env)
            except subprocess.TimeoutExpired:
                return f"Error: eksekusi melebihi batas {timeout} detik."
        result = f"exit_code={p.returncode}\n--- stdout ---\n{p.stdout}"
        if p.stderr:
            result += f"\n--- stderr ---\n{p.stderr}"
        return result

    def run_command(self, command: str, timeout: int = 60) -> str:
        timeout = max(1, min(int(timeout), 300))
        if not self.approvals.ask("Jalankan perintah shell", Syntax(command, "bash"), force=True):
            return "DITOLAK oleh pengguna."
        try:
            p = subprocess.run(
                command, shell=True, cwd=self.root, capture_output=True, text=True,
                timeout=timeout, env=self._clean_env(),
            )
        except subprocess.TimeoutExpired:
            return f"Error: perintah melebihi batas {timeout} detik."
        return f"exit_code={p.returncode}\n--- stdout ---\n{p.stdout}\n--- stderr ---\n{p.stderr}"

    # -- laporan --------------------------------------------------------------

    def create_report(self, filename: str, content: str, format: str = "md") -> str:
        fmt = format.lower()
        stem = Path(filename).stem
        rel = f"reports/{stem}.{'docx' if fmt == 'docx' else 'md'}"
        target = self._safe_path(rel)
        if fmt != "docx":
            return self._apply_write(target, rel, content)

        try:
            from docx import Document
        except ImportError:
            return "Error: python-docx belum terpasang (pip install python-docx). Gunakan format 'md'."
        if not self.approvals.ask(f"Buat laporan Word: {rel}", Syntax(content[:3000], "markdown")):
            return "DITOLAK oleh pengguna."
        doc = Document()
        for line in content.splitlines():
            if m := re.match(r"^(#{1,4})\s+(.*)", line):
                doc.add_heading(m.group(2), level=len(m.group(1)))
            elif re.match(r"^\s*[-*]\s+", line):
                doc.add_paragraph(re.sub(r"^\s*[-*]\s+", "", line), style="List Bullet")
            elif re.match(r"^\s*\d+\.\s+", line):
                doc.add_paragraph(re.sub(r"^\s*\d+\.\s+", "", line), style="List Number")
            elif line.strip():
                doc.add_paragraph(line)
        target.parent.mkdir(parents=True, exist_ok=True)
        doc.save(target)
        return f"OK: laporan disimpan di {rel}"


def _obj(props: dict[str, Any], required: list[str] | None = None) -> dict[str, Any]:
    schema: dict[str, Any] = {"type": "object", "properties": props}
    if required:
        schema["required"] = required
    return schema


TOOLS_SCHEMA = [
    {
        "name": "list_project_files",
        "description": "Daftar berkas kode dalam proyek (mengabaikan node_modules, .git, dll).",
        "parameters": _obj({
            "subdir": {"type": "string", "description": "Subdirektori relatif (default '.')"},
            "max_results": {"type": "integer", "description": "Maksimum berkas (default 300)"},
        }),
    },
    {
        "name": "search_in_files",
        "description": "Cari teks/regex di seluruh proyek. Kembalikan berkas:baris. Gunakan sebelum membaca berkas besar.",
        "parameters": _obj({
            "pattern": {"type": "string", "description": "Teks atau regex"},
            "file_pattern": {"type": "string", "description": "Filter glob nama berkas, mis. '*.py' (default '*')"},
            "context": {"type": "integer", "description": "Jumlah baris konteks di sekitar kecocokan (0-5)"},
        }, ["pattern"]),
    },
    {
        "name": "read_file_lines",
        "description": "Baca berkas pada rentang baris tertentu dengan nomor baris.",
        "parameters": _obj({
            "filepath": {"type": "string", "description": "Path relatif dari root proyek"},
            "start_line": {"type": "integer", "description": "Baris awal (1-indexed)"},
            "end_line": {"type": "integer", "description": "Baris akhir (default 400)"},
        }, ["filepath"]),
    },
    {
        "name": "write_file",
        "description": "Membuat atau menimpa seluruh isi berkas. Pengguna akan melihat diff dan harus menyetujui.",
        "parameters": _obj({
            "filepath": {"type": "string", "description": "Path relatif berkas"},
            "content": {"type": "string", "description": "Isi lengkap berkas"},
        }, ["filepath", "content"]),
    },
    {
        "name": "replace_in_file",
        "description": "Ganti satu potongan teks yang unik di berkas (perubahan minimal, lebih disukai daripada write_file).",
        "parameters": _obj({
            "filepath": {"type": "string"},
            "old_str": {"type": "string", "description": "Teks persis yang diganti; harus muncul tepat sekali"},
            "new_str": {"type": "string", "description": "Teks pengganti"},
        }, ["filepath", "old_str", "new_str"]),
    },
    {
        "name": "run_python",
        "description": (
            "Code interpreter: jalankan skrip Python mandiri untuk perhitungan, analisis data, atau verifikasi logika. "
            "Gunakan print() untuk output. Simpan grafik/CSV ke folder di variabel lingkungan AI_OUTPUT_DIR."
        ),
        "parameters": _obj({
            "code": {"type": "string", "description": "Kode Python lengkap"},
            "timeout": {"type": "integer", "description": "Batas detik (default 30, maks 180)"},
        }, ["code"]),
    },
    {
        "name": "run_command",
        "description": "Jalankan perintah shell di root proyek (tes, linter, build). Selalu butuh persetujuan pengguna.",
        "parameters": _obj({
            "command": {"type": "string"},
            "timeout": {"type": "integer", "description": "Batas detik (default 60, maks 300)"},
        }, ["command"]),
    },
    {
        "name": "create_report",
        "description": "Simpan laporan terstruktur ke folder reports/ sebagai Markdown atau Word (.docx).",
        "parameters": _obj({
            "filename": {"type": "string", "description": "Nama berkas tanpa ekstensi"},
            "content": {"type": "string", "description": "Isi laporan dalam Markdown"},
            "format": {"type": "string", "description": "'md' (default) atau 'docx'"},
        }, ["filename", "content"]),
    },
]


# ---------------------------------------------------------------------------
# Engine Multi-Provider: Streaming + Agentic Loop + Thinking + Multimodal
# ---------------------------------------------------------------------------

@dataclass
class AgentEngine:
    provider: str
    workspace: WorkspaceTools
    model: str = ""
    think: str = "medium"
    history: list[dict[str, str]] = field(default_factory=list)
    compacted: bool = False

    def __post_init__(self) -> None:
        self.provider = self.provider.lower().strip()
        self.think = self.think if self.think in THINK_LEVELS else "medium"
        if not self.model:
            self.model = DEFAULT_MODELS.get(self.provider, "")
        self._init_client()

    def _init_client(self) -> None:
        if self.provider == "openai":
            from openai import OpenAI
            self.client = OpenAI(api_key=os.environ["OPENAI_API_KEY"])
        elif self.provider == "anthropic":
            from anthropic import Anthropic
            self.client = Anthropic(api_key=os.environ["ANTHROPIC_API_KEY"])
        elif self.provider == "gemini":
            from google import genai
            self.client = genai.Client(api_key=os.environ["GEMINI_API_KEY"])
        else:
            raise ValueError(f"Provider tidak didukung: {self.provider}")

    def switch_provider(self, provider: str, model: str = "") -> None:
        self.provider, self.model = provider.lower().strip(), model
        self.model = model or DEFAULT_MODELS.get(self.provider, "")
        self._init_client()

    def _system(self) -> str:
        return f"{SYSTEM_PROMPT}\nWorkspace Root: {self.workspace.root}\nTanggal hari ini: {date.today().isoformat()}"

    # -- manajemen konteks ----------------------------------------------------

    def context_chars(self) -> int:
        return sum(len(m["content"]) for m in self.history)

    def _compact_history(self) -> None:
        """Ringkas pesan lama jadi satu ringkasan bila konteks melewati ambang."""
        if self.context_chars() <= MAX_HISTORY_CHARS or len(self.history) <= KEEP_RECENT_MESSAGES + 2:
            return
        old, keep = self.history[:-KEEP_RECENT_MESSAGES], self.history[-KEEP_RECENT_MESSAGES:]
        transcript = "\n\n".join(f"[{m['role']}]\n{m['content'][:8000]}" for m in old)[-250_000:]
        console.print("[dim]🗜  Konteks panjang — meringkas percakapan lama...[/dim]")
        try:
            summary = self._complete(
                "Anda meringkas percakapan teknis. Pertahankan: tujuan pengguna, keputusan, path berkas, nama fungsi, "
                "angka penting, bug yang ditemukan, perubahan yang sudah dibuat, dan hal yang belum selesai. "
                "Tulis ringkas dan faktual.",
                transcript,
            )
            self.history = [
                {"role": "user", "content": f"[Ringkasan percakapan sebelumnya]\n{summary}"},
                {"role": "assistant", "content": "Dipahami. Saya melanjutkan dengan konteks tersebut."},
                *keep,
            ]
        except Exception as e:  # noqa: BLE001
            console.print(f"[yellow]Ringkasan gagal ({e}); memangkas pesan lama.[/yellow]")
            while self.context_chars() > MAX_HISTORY_CHARS and len(self.history) > KEEP_RECENT_MESSAGES:
                self.history.pop(0)
            while self.history and self.history[0]["role"] != "user":
                self.history.pop(0)
        self.compacted = True

    def _complete(self, system: str, user: str) -> str:
        """Pemanggilan non-streaming sederhana (untuk ringkasan)."""
        if self.provider == "openai":
            r = with_retry(lambda: self.client.chat.completions.create(
                model=self.model,
                messages=[{"role": "system", "content": system}, {"role": "user", "content": user}],
            ))
            return r.choices[0].message.content or ""
        if self.provider == "anthropic":
            r = with_retry(lambda: self.client.messages.create(
                model=self.model, max_tokens=4096, system=system,
                messages=[{"role": "user", "content": user}],
            ))
            return "".join(b.text for b in r.content if b.type == "text")
        from google.genai import types
        r = with_retry(lambda: self.client.models.generate_content(
            model=self.model, contents=user,
            config=types.GenerateContentConfig(system_instruction=system),
        ))
        return r.text or ""

    # -- entry point ----------------------------------------------------------

    def ask_stream(self, user_prompt: str, attachments: list[Attachment] | None = None) -> Generator[str, None, str]:
        attachments = attachments or []
        unsupported = [a for a in attachments if a.kind not in PROVIDER_MEDIA[self.provider]]
        if unsupported:
            names = ", ".join(f"{a.path.name} ({a.kind})" for a in unsupported)
            raise ValueError(
                f"Provider '{self.provider}' tidak mendukung: {names}. "
                f"Gunakan /provider gemini untuk audio/video."
            )

        self.history.append({"role": "user", "content": user_prompt})
        try:
            self._compact_history()
            runner = {"openai": self._run_openai, "anthropic": self._run_anthropic, "gemini": self._run_gemini}[
                self.provider
            ]
            answer = yield from runner(attachments)
        except BaseException:
            # Jangan tinggalkan pesan user tanpa jawaban di riwayat.
            if self.history and self.history[-1]["role"] == "user" and self.history[-1]["content"] == user_prompt:
                self.history.pop()
            raise
        self.history.append({"role": "assistant", "content": answer})
        return answer

    def _call_tool(self, name: str, args: dict[str, Any]) -> str:
        shown = json.dumps(args, ensure_ascii=False)
        console.print(f"[dim cyan]🔧 {name}({truncate(shown, 160)})[/dim cyan]")
        return self.workspace.execute(name, args)

    # -- OpenAI ---------------------------------------------------------------

    def _openai_is_reasoning(self) -> bool:
        return self.model.startswith(("o1", "o3", "o4", "gpt-5"))

    @staticmethod
    def _oai_user_content(text: str, atts: list[Attachment]) -> Any:
        if not atts:
            return text
        parts: list[dict[str, Any]] = [{"type": "text", "text": text}]
        for a in atts:
            data_url = f"data:{a.mime};base64,{a.b64}"
            if a.kind == "image":
                parts.append({"type": "image_url", "image_url": {"url": data_url}})
            else:
                parts.append({"type": "file", "file": {"filename": a.path.name, "file_data": data_url}})
        return parts

    def _run_openai(self, atts: list[Attachment]) -> Generator[str, None, str]:
        tools = [{"type": "function", "function": t} for t in TOOLS_SCHEMA]
        hist = [dict(m) for m in self.history]
        hist[-1]["content"] = self._oai_user_content(hist[-1]["content"], atts)
        messages: list[dict[str, Any]] = [{"role": "system", "content": self._system()}, *hist]
        reasoning = self._openai_is_reasoning()
        out: list[str] = []

        for _ in range(MAX_AGENT_STEPS):
            kwargs: dict[str, Any] = dict(model=self.model, messages=messages, tools=tools, stream=True)
            if reasoning:
                kwargs["reasoning_effort"] = "low" if self.think == "off" else self.think
            else:
                kwargs["temperature"] = 0.2
            stream = with_retry(lambda: self.client.chat.completions.create(**kwargs))

            step_text: list[str] = []
            calls: dict[int, dict[str, str]] = {}
            for chunk in stream:
                if not chunk.choices:
                    continue
                delta = chunk.choices[0].delta
                if delta.content:
                    step_text.append(delta.content)
                    out.append(delta.content)
                    yield delta.content
                for tc in delta.tool_calls or []:
                    c = calls.setdefault(tc.index, {"id": "", "name": "", "args": ""})
                    if tc.id:
                        c["id"] = tc.id
                    if tc.function:
                        if tc.function.name:
                            c["name"] = tc.function.name
                        if tc.function.arguments:
                            c["args"] += tc.function.arguments

            if not calls:
                return "".join(out)

            messages.append({
                "role": "assistant",
                "content": "".join(step_text) or None,
                "tool_calls": [
                    {"id": c["id"], "type": "function", "function": {"name": c["name"], "arguments": c["args"] or "{}"}}
                    for _, c in sorted(calls.items())
                ],
            })
            for _, c in sorted(calls.items()):
                try:
                    args = json.loads(c["args"] or "{}")
                except json.JSONDecodeError as e:
                    result = f"Error: argumen JSON tidak valid: {e}"
                else:
                    result = self._call_tool(c["name"], args)
                messages.append({"role": "tool", "tool_call_id": c["id"], "content": result})

        note = "\n\n⚠️ Batas langkah agen tercapai. Lanjutkan dengan pesan berikutnya bila perlu."
        yield note
        return "".join(out) + note

    # -- Anthropic ------------------------------------------------------------

    @staticmethod
    def _anth_user_content(text: str, atts: list[Attachment]) -> Any:
        if not atts:
            return text
        blocks: list[dict[str, Any]] = [
            {
                "type": "image" if a.kind == "image" else "document",
                "source": {"type": "base64", "media_type": a.mime, "data": a.b64},
            }
            for a in atts
        ]
        blocks.append({"type": "text", "text": text})
        return blocks

    def _run_anthropic(self, atts: list[Attachment]) -> Generator[str, None, str]:
        tools = [{"name": t["name"], "description": t["description"], "input_schema": t["parameters"]} for t in TOOLS_SCHEMA]
        messages: list[dict[str, Any]] = [dict(m) for m in self.history]
        messages[-1]["content"] = self._anth_user_content(messages[-1]["content"], atts)
        system = [{"type": "text", "text": self._system(), "cache_control": {"type": "ephemeral"}}]
        budget = ANTHROPIC_BUDGET[self.think]
        out: list[str] = []

        for _ in range(MAX_AGENT_STEPS):
            kwargs: dict[str, Any] = dict(
                model=self.model,
                max_tokens=MAX_OUTPUT_TOKENS + budget,
                system=system,
                messages=messages,
                tools=tools,
            )
            if budget:
                kwargs["thinking"] = {"type": "enabled", "budget_tokens": budget}
            else:
                kwargs["temperature"] = 0.2

            mgr = self.client.messages.stream(**kwargs)
            stream = with_retry(mgr.__enter__)
            try:
                for text in stream.text_stream:
                    out.append(text)
                    yield text
                final = stream.get_final_message()
            finally:
                mgr.__exit__(None, None, None)

            tool_uses = [b for b in final.content if b.type == "tool_use"]
            if not tool_uses or final.stop_reason != "tool_use":
                if final.stop_reason == "max_tokens":
                    note = "\n\n⚠️ Jawaban terpotong (batas token). Minta saya melanjutkan."
                    out.append(note)
                    yield note
                return "".join(out)

            messages.append({"role": "assistant", "content": final.content})
            results = [
                {"type": "tool_result", "tool_use_id": tu.id, "content": self._call_tool(tu.name, dict(tu.input))}
                for tu in tool_uses
            ]
            messages.append({"role": "user", "content": results})

        note = "\n\n⚠️ Batas langkah agen tercapai. Lanjutkan dengan pesan berikutnya bila perlu."
        yield note
        return "".join(out) + note

    # -- Gemini ---------------------------------------------------------------

    def _gemini_part(self, a: Attachment):
        from google.genai import types
        if a.kind in ("audio", "video") or a.path.stat().st_size > 15_000_000:
            console.print(f"[dim]⬆️  Mengunggah {a.path.name}...[/dim]")
            f = with_retry(lambda: self.client.files.upload(file=str(a.path)))
            while f.state is not None and f.state.name == "PROCESSING":
                time.sleep(2)
                f = self.client.files.get(name=f.name)
            if f.state is not None and f.state.name == "FAILED":
                raise RuntimeError(f"Gemini gagal memproses {a.path.name}")
            return types.Part.from_uri(file_uri=f.uri, mime_type=f.mime_type or a.mime)
        return types.Part.from_bytes(data=a.path.read_bytes(), mime_type=a.mime)

    def _run_gemini(self, atts: list[Attachment]) -> Generator[str, None, str]:
        from google.genai import types

        decls = [
            types.FunctionDeclaration(name=t["name"], description=t["description"], parameters=t["parameters"])
            for t in TOOLS_SCHEMA
        ]
        contents: list[Any] = [
            types.Content(
                role="user" if m["role"] == "user" else "model",
                parts=[types.Part.from_text(text=m["content"])],
            )
            for m in self.history
        ]
        if atts:
            contents[-1].parts.extend(self._gemini_part(a) for a in atts)

        config = types.GenerateContentConfig(
            system_instruction=self._system(),
            tools=[types.Tool(function_declarations=decls)],
            temperature=0.2,
            thinking_config=types.ThinkingConfig(thinking_budget=GEMINI_BUDGET[self.think]),
            automatic_function_calling=types.AutomaticFunctionCallingConfig(disable=True),
        )
        out: list[str] = []

        for _ in range(MAX_AGENT_STEPS):
            parts: list[Any] = []
            stream = with_retry(
                lambda: self.client.models.generate_content_stream(model=self.model, contents=contents, config=config)
            )
            for chunk in stream:
                for cand in chunk.candidates or []:
                    if not cand.content or not cand.content.parts:
                        continue
                    for p in cand.content.parts:
                        parts.append(p)
                        if p.text and not getattr(p, "thought", False):
                            out.append(p.text)
                            yield p.text

            calls = [p.function_call for p in parts if p.function_call]
            if not calls:
                return "".join(out)

            contents.append(types.Content(role="model", parts=parts))
            responses = [
                types.Part.from_function_response(
                    name=fc.name, response={"result": self._call_tool(fc.name, dict(fc.args or {}))}
                )
                for fc in calls
            ]
            contents.append(types.Content(role="user", parts=responses))

        note = "\n\n⚠️ Batas langkah agen tercapai. Lanjutkan dengan pesan berikutnya bila perlu."
        yield note
        return "".join(out) + note


# ---------------------------------------------------------------------------
# Antarmuka Command-Line (CLI)
# ---------------------------------------------------------------------------

def print_help() -> None:
    t = Table(title="Perintah AI Coder", border_style="cyan")
    t.add_column("Perintah", style="yellow", no_wrap=True)
    t.add_column("Fungsi")
    rows = [
        ("/workspace <path>", "Ubah direktori proyek"),
        ("/file <path>", "Sematkan isi berkas teks/kode ke pesan berikutnya"),
        ("/attach <path>", "Lampirkan gambar/PDF/audio/video ke pesan berikutnya (audio/video: Gemini)"),
        ("/think <off|low|medium|high>", "Atur kedalaman penalaran (deep thinking)"),
        ("/provider <openai|anthropic|gemini> [model]", "Ganti provider/model di tengah sesi"),
        ("/model <nama>", "Ganti model pada provider aktif"),
        ("/approve", "Aktif/nonaktifkan persetujuan otomatis untuk perubahan berkas"),
        ("/undo", "Batalkan perubahan berkas terakhir oleh AI"),
        ("/session <nama>", "Beralih/muat sesi"),
        ("/sessions", "Daftar sesi tersimpan"),
        ("/status", "Info provider, model, konteks"),
        ("/reset", "Hapus riwayat sesi aktif"),
        ("/help", "Bantuan"),
        ("/exit", "Keluar"),
    ]
    for r in rows:
        t.add_row(*r)
    console.print(t)


def main() -> int:
    session_id = "default"
    approvals = Approvals()
    workspace = WorkspaceTools(Path.cwd(), approvals)
    store = SessionStore()

    try:
        engine = AgentEngine(
            provider=os.getenv("PROVIDER", "anthropic"),
            model=os.getenv("MODEL", ""),
            think=os.getenv("THINK", "medium"),
            workspace=workspace,
        )
    except KeyError as e:
        console.print(f"[bold red]API Key belum diatur di environment: {e}[/bold red]")
        return 1
    except Exception as e:  # noqa: BLE001
        console.print(f"[bold red]Gagal menginisialisasi engine: {e}[/bold red]")
        return 1

    engine.history = store.load_session(session_id)
    live_holder: dict[str, Live | None] = {"live": None}
    approvals.before = lambda: live_holder["live"] and live_holder["live"].stop()
    approvals.after = lambda: live_holder["live"] and live_holder["live"].start()

    def banner() -> Panel:
        return Panel.fit(
            f"[bold cyan]AI Coder v2 (Agentic)[/bold cyan]\n"
            f"Provider : [green]{engine.provider}[/green] ({engine.model})  |  Thinking: [green]{engine.think}[/green]\n"
            f"Workspace: [yellow]{workspace.root}[/yellow]\n"
            f"Sandbox  : {SANDBOX}  |  Sesi: [magenta]{session_id}[/magenta] ({len(engine.history)} pesan)\n"
            f"Ketik [yellow]/help[/yellow] untuk daftar perintah.",
            border_style="cyan",
        )

    console.print(banner())
    pinned: str | None = None
    pending: list[Attachment] = []

    while True:
        try:
            user_input = console.input(f"\n[bold blue]({session_id}) Anda >[/bold blue] ").strip()
        except (EOFError, KeyboardInterrupt):
            console.print("\n[dim]Sampai jumpa.[/dim]")
            break
        if not user_input:
            continue

        if user_input.startswith("/"):
            cmd, _, arg = user_input.partition(" ")
            cmd, arg = cmd.lower(), arg.strip()
            try:
                if cmd in ("/exit", "/quit"):
                    break
                elif cmd == "/help":
                    print_help()
                elif cmd == "/status":
                    console.print(banner())
                    console.print(
                        f"Konteks: ~{engine.context_chars():,} karakter dari ambang {MAX_HISTORY_CHARS:,} | "
                        f"Auto-approve: {approvals.auto} | Lampiran tertunda: {len(pending)}"
                    )
                elif cmd == "/reset":
                    engine.history.clear()
                    store.clear_session(session_id)
                    console.print("[dim]Riwayat sesi ini dihapus.[/dim]")
                elif cmd == "/workspace":
                    if arg:
                        workspace.set_root(Path(arg))
                    console.print(f"Workspace: [yellow]{workspace.root}[/yellow]")
                elif cmd == "/file":
                    if not arg:
                        console.print("[red]Contoh: /file src/main.py[/red]")
                    else:
                        pinned = workspace.read_file_lines(arg, 1, 800)
                        console.print(f"[green]'{arg}' disematkan untuk pesan berikutnya.[/green]")
                elif cmd == "/attach":
                    if not arg:
                        console.print("[red]Contoh: /attach diagram.png[/red]")
                    else:
                        att = load_attachment(arg, workspace.root)
                        if att.kind not in PROVIDER_MEDIA[engine.provider]:
                            console.print(f"[yellow]Catatan: provider '{engine.provider}' belum mendukung {att.kind}; gunakan /provider gemini.[/yellow]")
                        pending.append(att)
                        console.print(f"[green]Terlampir: {att.path.name} ({att.kind})[/green]")
                elif cmd == "/think":
                    if arg not in THINK_LEVELS:
                        console.print(f"[red]Pilih salah satu: {', '.join(THINK_LEVELS)}[/red]")
                    else:
                        engine.think = arg
                        console.print(f"[green]Thinking: {arg}[/green]")
                elif cmd == "/provider":
                    name, _, mdl = arg.partition(" ")
                    engine.switch_provider(name, mdl.strip())
                    console.print(f"[green]Provider: {engine.provider} ({engine.model})[/green]")
                elif cmd == "/model":
                    if arg:
                        engine.model = arg
                    console.print(f"Model: [green]{engine.model}[/green]")
                elif cmd == "/approve":
                    approvals.auto = not approvals.auto
                    console.print(f"Auto-approve perubahan berkas: [bold]{approvals.auto}[/bold] (perintah shell/kode lokal tetap minta izin)")
                elif cmd == "/undo":
                    console.print(workspace.undo())
                elif cmd == "/session":
                    if arg:
                        session_id = arg
                        engine.history = store.load_session(session_id)
                    console.print(f"Sesi: [magenta]{session_id}[/magenta] ({len(engine.history)} pesan)")
                elif cmd == "/sessions":
                    rows = store.list_sessions()
                    if not rows:
                        console.print("[dim]Belum ada sesi tersimpan.[/dim]")
                    for sid, count, updated in rows:
                        console.print(f"• [magenta]{sid}[/magenta] — {count} pesan (terakhir: {updated})")
                else:
                    console.print(f"[red]Perintah tidak dikenal: {cmd}. Ketik /help.[/red]")
            except KeyError as e:
                console.print(f"[red]API key belum diatur: {e}[/red]")
            except Exception as e:  # noqa: BLE001
                console.print(f"[red]{e}[/red]")
            continue

        # -- siapkan prompt --
        full_prompt = user_input
        if pinned:
            full_prompt = f"Berikut isi berkas yang saya sematkan:\n\n{pinned}\n\n---\n{full_prompt}"
            pinned = None
        if pending:
            full_prompt += "\n\n[Lampiran: " + ", ".join(a.path.name for a in pending) + "]"
        attachments, pending = pending, []

        # -- eksekusi + live rendering --
        accumulated = ""
        try:
            live = Live(
                Panel("[dim]berpikir...[/dim]", title="AI", border_style="cyan"),
                console=console, refresh_per_second=10, vertical_overflow="visible",
            )
            live_holder["live"] = live
            with live:
                for chunk in engine.ask_stream(full_prompt, attachments):
                    accumulated += chunk
                    live.update(Panel(Markdown(accumulated), title=f"AI ({engine.model})", border_style="cyan"))
            store.save_message(session_id, "user", full_prompt)
            store.save_message(session_id, "assistant", accumulated)
            if engine.compacted:  # riwayat dikompaksi: sinkronkan ke database
                store.replace_session(session_id, engine.history)
                engine.compacted = False
        except KeyboardInterrupt:
            console.print("\n[yellow]Dihentikan pengguna.[/yellow]")
        except Exception as e:  # noqa: BLE001
            console.print(f"[bold red]Terjadi kesalahan: {e}[/bold red]")
        finally:
            live_holder["live"] = None

    return 0


if __name__ == "__main__":
    sys.exit(main())
